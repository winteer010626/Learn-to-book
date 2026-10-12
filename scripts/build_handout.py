# -*- coding: utf-8 -*-
"""
build_handout.py — 数据驱动的 Word 讲义生成器。
用法:
    python build_handout.py spec.json out.docx
spec 为 UTF-8 JSON，结构见同 skill 的 references/spec_example.json。
所有排版常量（深蓝表头、边框、字体、A4、页脚 PAGE 域）已内置并通过 audit。
"""
import sys, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- 排版常量（如无用户特别要求，不要改） ----------
HEADER_FILL = "1F3A5F"   # 表头深蓝
BORDER_COLOR = "4F81BD"  # 表格边框蓝
HEAD_COLOR  = RGBColor(0x1F, 0x3A, 0x5F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ASCII_FONT = "Calibri"
EA_FONT = "宋体"
HEAD_EA_FONT = "微软雅黑"
BODY_SIZE = 10.5
TABLE_SIZE = 10

ALIGN = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
         "r": WD_ALIGN_PARAGRAPH.RIGHT, "j": WD_ALIGN_PARAGRAPH.JUSTIFY}


def set_font(run, size=BODY_SIZE, bold=False, italic=False, color=None,
             ascii_font=ASCII_FONT, ea_font=EA_FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), ea_font)
    rfonts.set(qn("w:cs"), ascii_font)


def setup_document():
    doc = Document()
    # Normal 样式
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(4)
    # 页面：A4，四边 2.5cm
    for sec in doc.sections:
        sec.page_width = Cm(21)
        sec.page_height = Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.left_margin = sec.right_margin = Cm(2.5)
    # 标题样式颜色
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = ASCII_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = HEAD_COLOR
        st.element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_EA_FONT)
    add_footer_page(doc)
    return doc


def add_footer_page(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    set_font(run, size=9)


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_font(r, size=18, bold=True, color=HEAD_COLOR, ea_font=HEAD_EA_FONT)
    p.paragraph_format.space_after = Pt(4)
    if subtitle:
        ps = doc.add_paragraph()
        ps.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = ps.add_run(subtitle)
        set_font(rs, size=10.5, italic=True)
        ps.paragraph_format.space_after = Pt(10)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading %d" % level)
    r = p.add_run(text)
    set_font(r, size=14 if level == 1 else 12, bold=True,
             color=HEAD_COLOR, ea_font=HEAD_EA_FONT)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5 if level == 1 else 4)
    p.paragraph_format.keep_with_next = True
    return p


def add_para(doc, block):
    p = doc.add_paragraph()
    p.alignment = ALIGN.get(block.get("align", "l"), WD_ALIGN_PARAGRAPH.LEFT)
    if block.get("space_before"):
        p.paragraph_format.space_before = Pt(block["space_before"])
    if "spans" in block:
        for sp in block["spans"]:
            r = p.add_run(sp.get("text", ""))
            set_font(r, size=sp.get("size", block.get("size", BODY_SIZE)),
                     bold=sp.get("bold", False), italic=sp.get("italic", False))
    else:
        r = p.add_run(block.get("text", ""))
        set_font(r, size=block.get("size", BODY_SIZE),
                 bold=block.get("bold", False), italic=block.get("italic", False))
    return p


def add_bullet(doc, block):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(block.get("text", ""))
    set_font(r, size=block.get("size", BODY_SIZE),
             bold=block.get("bold", False), italic=block.get("italic", False))
    p.paragraph_format.space_after = Pt(2)
    return p


def _set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
    trPr.append(th)
    cs = OxmlElement("w:cantSplit"); trPr.append(cs)


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


# tblPr 相关节点的 schema 顺序（值越小越靠前）
_TBLPR_RANK = {qn("w:tblW"): 0, qn("w:jc"): 1, qn("w:tblBorders"): 2,
               qn("w:tblLayout"): 3, qn("w:tblLook"): 4}


def _tblpr_insert(tblPr, element):
    rank = _TBLPR_RANK[element.tag]
    for child in tblPr:
        if child.tag in _TBLPR_RANK and _TBLPR_RANK[child.tag] > rank:
            child.addprevious(element)
            return
    tblPr.append(element)


def _set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), BORDER_COLOR)
        borders.append(e)
    _tblpr_insert(tblPr, borders)


def _shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    va = tcPr.find(qn("w:vAlign"))
    if va is not None:
        va.addprevious(sh)
    else:
        tcPr.append(sh)


def _fill_cell(cell, text, align, size, bold, color, ea_font):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("" if text is None else str(text))
    set_font(r, size=size, bold=bold, color=color, ea_font=ea_font)


def add_table(doc, block):
    headers = block.get("headers", [])
    rows = block.get("rows", [])
    ncol = len(headers) if headers else (len(rows[0]) if rows else 1)
    aligns = block.get("aligns", ["l"] * ncol)
    widths = block.get("widths", [])
    fill = block.get("header_fill", HEADER_FILL)

    tbl = doc.add_table(rows=1, cols=ncol)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 固定布局（按 schema 顺序插入，勿直接 append）
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    _tblpr_insert(tblPr, layout)
    _set_borders(tbl)

    hdr = tbl.rows[0]
    _set_repeat_header(hdr)
    for j, h in enumerate(headers):
        _fill_cell(hdr.cells[j], h, "c", TABLE_SIZE, True, WHITE, HEAD_EA_FONT)
        _shade_cell(hdr.cells[j], fill)
    for r_i, rowdata in enumerate(rows):
        row = tbl.add_row()
        _cant_split(row)
        for j in range(ncol):
            val = rowdata[j] if j < len(rowdata) else ""
            _fill_cell(row.cells[j], val, aligns[j] if j < len(aligns) else "l",
                       TABLE_SIZE, False, None, EA_FONT)
    # 列宽（逐格设置，固定布局生效）
    if widths:
        for j, w in enumerate(widths):
            for row in tbl.rows:
                row.cells[j].width = Cm(float(w))
    if block.get("note"):
        p = doc.add_paragraph()
        r = p.add_run(block["note"])
        set_font(r, size=9.5, italic=True)
        p.paragraph_format.space_after = Pt(4)
    return tbl


def build(spec, out_path):
    doc = setup_document()
    add_title(doc, spec["title"], spec.get("subtitle"))
    for block in spec.get("blocks", []):
        t = block.get("t")
        if t == "h1":
            add_heading(doc, block["text"], 1)
        elif t == "h2":
            add_heading(doc, block["text"], 2)
        elif t == "p" or t == "spans":
            add_para(doc, block)
        elif t == "bullet":
            add_bullet(doc, block)
        elif t == "table":
            add_table(doc, block)
        elif t == "pagebreak":
            doc.add_page_break()
    doc.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_handout.py spec.json out.docx")
        sys.exit(1)
    with open(sys.argv[1], "r", encoding="utf-8-sig") as f:
        spec = json.load(f)
    build(spec, sys.argv[2])
    print("SAVED", sys.argv[2])
